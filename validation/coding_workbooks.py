#!/usr/bin/env python3
"""Excel workbooks for the two blinded coders, and their merge into a gold-standard draft.

Three steps, each a sub-command (requires ``openpyxl``; see requirements.txt):

1. ``enumerate``  -> case_enumeration.xlsx
   One sheet listing the locked PRIMARY studies with empty rows. The person
   who enumerates (not one of the coders' decisions) lists every candidate
   effect of each paper by LOCATION only - page, table, model, row/column -
   including non-focal coefficients (controls, interactions, other outcomes),
   so precision and recall have real denominators. No values, no in-scope
   judgement are written here.

2. ``coders``     -> coder1.xlsx, coder2.xlsx
   From the filled enumeration: identical workbooks with one row per
   case_id = <paper_id>-<effect_id>. Identifier columns are locked; input
   cells carry drop-down lists and range checks (r in [-1, 1], N > 0).
   Each coder works on their own file, never sees the other file or any
   machine proposal before both are marked complete.

3. ``merge``      -> gold_standard_draft.csv, disagreements.csv
   Reads both completed workbooks. Where the coders agree on scope, route,
   r (within the tolerance), N and both measure classes, the gold_* columns
   are pre-filled and marked "agreed"; every other case is left blank in
   gold_* and listed in disagreements.csv for adjudication against the full
   text. manual_minutes is coder 1's time (coder 1 is the PI, the same person
   who later times the assisted verification, so the time comparison is
   within-person). The two coders' own columns are never edited.

    python validation/coding_workbooks.py enumerate --lock validation/sampling/sample_lock_v1.json --out validation/coding/case_enumeration.xlsx
    python validation/coding_workbooks.py coders --enumeration validation/coding/case_enumeration.xlsx --out-dir validation/coding
    python validation/coding_workbooks.py merge --coder1 validation/coding/coder1.xlsx --coder2 validation/coding/coder2.xlsx --out-dir validation/coding
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SAMPLING = ROOT / "validation" / "sampling" / "candidate_sample_v1.csv"

ROUTES = ["direct_r", "t_to_r", "beta_to_r", "other"]
DOI = ["FSTS", "GEO", "EXP", "FDI", "COMP", "OTH"]
PERF = ["ACC", "MKT", "LAB", "MIX"]
BOOL = ["TRUE", "FALSE"]
CODING_COLUMNS = ["case_id", "paper_id", "effect_id", "source_locator",
                  "in_scope", "route", "r", "n", "doi_measure", "performance_measure", "minutes", "notes"]
GOLD_COLUMNS = ["case_id", "paper_id", "effect_id", "source_locator", "coder1_in_scope", "coder2_in_scope",
                "gold_in_scope", "coder1_route", "coder2_route", "gold_route", "gold_r", "gold_n",
                "gold_doi_measure", "gold_performance_measure", "manual_minutes", "adjudication_notes"]

GUIDE_ENUM = [
    "LIỆT KÊ CANDIDATE EFFECT (chỉ vị trí, không giá trị)",
    "Mỗi bài: liệt kê MỌI hệ số có thể bị nhầm là hệ số quốc tế hóa → hiệu quả, kể cả biến kiểm soát, tương tác, mẫu phụ, kết quả không trọng tâm.",
    "effect_id: E1, E2, ... trong từng bài. source_locator: trang + bảng + mô hình + dòng/cột, ví dụ 'p.9 Table 3 Model 2, row FSTS, col ROA'.",
    "KHÔNG ghi giá trị r, t, β hay N; KHÔNG ghi hệ số nào là trọng tâm. Hai người mã hóa sẽ tự quyết.",
    "Thêm dòng nếu cần (giữ đúng paper_id). Bài nào cũng phải có ít nhất một dòng.",
]
GUIDE_CODER = [
    "MÃ HÓA ĐỘC LẬP (xem validation/CODING_MANUAL_VI.md)",
    "Không mở tệp của người mã hóa kia, không xem đề xuất của M-AIDA cho đến khi cả hai đã đánh dấu hoàn tất.",
    "in_scope: TRUE nếu candidate effect là hệ số I → P đủ điều kiện theo tiêu chí P6; FALSE nếu không.",
    "route: direct_r (bài báo cáo r), t_to_r (từ t và df), beta_to_r (từ β chuẩn hóa), other (tuyến khác có trong protocol). Để trống khi in_scope = FALSE.",
    "r: hệ số tương quan sau quy đổi, trong [-1, 1], 3 chữ số thập phân. n: cỡ mẫu dùng cho hệ số đó.",
    "minutes: số phút mã hóa candidate effect này (đọc, tìm, tính). notes: chỗ khó, giả định.",
    "Khi xong: điền coder_id và completed = TRUE ở trang 'Coder', lưu tệp, gửi cho người phân xử.",
]


def _need_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("openpyxl is required: pip install -r validation/requirements.txt") from exc


def _papers(lock: Path | None, sampling: Path) -> list[dict[str, str]]:
    with sampling.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if lock is not None:
        info = json.loads(lock.read_text(encoding="utf-8"))
        if info.get("status") != "LOCKED":
            raise SystemExit(f"{lock} is not a LOCKED sampling frame")
        wanted = list(info["primary_study_ids"])
    else:
        wanted = [r["study_id"] for r in rows if r.get("selection_group") == "PRIMARY"]
    by_id = {r["study_id"]: r for r in rows}
    return [by_id[sid] for sid in wanted]


def _guide(ws, lines: list[str]) -> None:
    from openpyxl.styles import Alignment, Font

    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=13)


def _header(ws, columns: list[str], widths: dict[str, int]) -> None:
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill("solid", fgColor="EFE8DA")
    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = fill
        ws.column_dimensions[cell.column_letter].width = widths.get(name, 14)
    ws.freeze_panes = "A2"


def cmd_enumerate(args: argparse.Namespace) -> int:
    _need_openpyxl()
    from openpyxl import Workbook

    papers = _papers(args.lock, args.sampling)
    if args.out.exists():
        raise SystemExit(f"{args.out} exists; outputs are never overwritten")
    wb = Workbook()
    _guide(wb.active, GUIDE_ENUM)
    wb.active.title = "Huong dan"
    ws = wb.create_sheet("Cases")
    cols = ["paper_id", "label", "year", "effect_id", "source_locator", "enumerator_notes"]
    _header(ws, cols, {"label": 34, "source_locator": 48, "enumerator_notes": 40, "effect_id": 10})
    r = 2
    for p in papers:
        for _ in range(args.rows_per_paper):
            ws.cell(row=r, column=1, value=p["study_id"])
            ws.cell(row=r, column=2, value=p.get("label", ""))
            ws.cell(row=r, column=3, value=int(p.get("year") or 0) or None)
            r += 1
    meta = wb.create_sheet("Meta")
    meta["A1"], meta["B1"] = "enumerator_id", ""
    meta["A2"], meta["B2"] = "lock_file", str(args.lock) if args.lock else "(unlocked: rehearsal only)"
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out}: {len(papers)} studies, {args.rows_per_paper} blank rows each")
    return 0


def read_enumeration(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["Cases"]
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) for name in ("paper_id", "effect_id", "source_locator")}
    cases, seen, problems = [], set(), []
    papers_seen: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        pid = str(row[idx["paper_id"]] or "").strip()
        if pid:
            papers_seen.add(pid)
        eid = str(row[idx["effect_id"]] or "").strip()
        loc = str(row[idx["source_locator"]] or "").strip()
        if not pid or (not eid and not loc):
            continue
        if not eid or not loc:
            problems.append(f"{pid}: effect_id and source_locator are both required ({eid or '?'} / {loc or '?'})")
            continue
        key = (pid, eid)
        if key in seen:
            problems.append(f"{pid}-{eid}: duplicated")
            continue
        seen.add(key)
        cases.append({"case_id": f"{pid}-{eid}", "paper_id": pid, "effect_id": eid, "source_locator": loc})
    with_cases = {c["paper_id"] for c in cases}
    for pid in sorted(papers_seen - with_cases):
        problems.append(f"{pid}: no candidate effect listed")
    if problems:
        raise SystemExit("enumeration is not complete:\n- " + "\n- ".join(problems))
    return cases


def cmd_coders(args: argparse.Namespace) -> int:
    _need_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Protection
    from openpyxl.worksheet.datavalidation import DataValidation

    cases = read_enumeration(args.enumeration)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    for coder in ("coder1", "coder2"):
        out = args.out_dir / f"{coder}.xlsx"
        if out.exists():
            raise SystemExit(f"{out} exists; outputs are never overwritten")
        wb = Workbook()
        _guide(wb.active, GUIDE_CODER)
        wb.active.title = "Huong dan"
        info = wb.create_sheet("Coder")
        info["A1"], info["A2"], info["A3"] = "coder_slot", "coder_id", "completed"
        info["B1"], info["B3"] = coder, "FALSE"
        for ref in ("B2", "B3"):
            info[ref].protection = Protection(locked=False)
        dv_done = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        info.add_data_validation(dv_done)
        dv_done.add("B3")
        info.protection.sheet = True
        ws = wb.create_sheet("Coding")
        _header(ws, CODING_COLUMNS, {"case_id": 14, "source_locator": 46, "notes": 40,
                                     "doi_measure": 14, "performance_measure": 21})
        n = len(cases) + 1
        from openpyxl.styles import PatternFill

        locked_fill = PatternFill("solid", fgColor="F2F2F2")
        for r, case in enumerate(cases, start=2):
            for c, key in enumerate(("case_id", "paper_id", "effect_id", "source_locator"), start=1):
                cell = ws.cell(row=r, column=c, value=case[key])
                cell.fill = locked_fill  # grey = locked identifier, white = the coder's input
            for c in range(5, len(CODING_COLUMNS) + 1):
                ws.cell(row=r, column=c).protection = Protection(locked=False)
        def lst(values: list[str], col: str, ws=ws, n=n) -> None:
            dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Giá trị không hợp lệ", error="Chọn từ danh sách.")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")
        lst(BOOL, "E")
        lst(ROUTES, "F")
        dv_r = DataValidation(type="decimal", operator="between", formula1="-1", formula2="1", allow_blank=True,
                              showErrorMessage=True, error="r phải nằm trong [-1, 1].")
        dv_n = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True,
                              showErrorMessage=True, error="n phải là số nguyên dương.")
        dv_m = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              showErrorMessage=True, error="Số phút không âm.")
        for dv, col in ((dv_r, "G"), (dv_n, "H"), (dv_m, "K")):
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")
        lst(DOI, "I")
        lst(PERF, "J")
        for r in range(2, n + 1):
            ws.cell(row=r, column=7).number_format = "0.000"
        ws.protection.sheet = True
        ws.protection.formatColumns = False
        wb.save(out)
        print(f"wrote {out}: {len(cases)} candidate effects")
    return 0


def _read_coder(path: Path) -> tuple[str, bool, dict[str, dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    info = wb["Coder"]
    coder_id = str(info["B2"].value or "").strip()
    completed = str(info["B3"].value).strip().upper() == "TRUE"
    ws = wb["Coding"]
    header = [c.value for c in ws[1]]
    if header[: len(CODING_COLUMNS)] != CODING_COLUMNS:
        raise SystemExit(f"{path}: unexpected columns {header}")
    rows: dict[str, dict[str, Any]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(CODING_COLUMNS, values, strict=False))  # extra columns ignored
        cid = str(rec["case_id"] or "").strip()
        if cid:
            rows[cid] = rec
    return coder_id, completed, rows


def _bool(value: Any, where: str) -> bool:
    text = str(value).strip().upper() if value is not None else ""
    if text in ("TRUE", "1"):
        return True
    if text in ("FALSE", "0"):
        return False
    raise SystemExit(f"{where}: in_scope must be TRUE or FALSE")


def _num(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def cmd_merge(args: argparse.Namespace) -> int:
    _need_openpyxl()
    id1, done1, c1 = _read_coder(args.coder1)
    id2, done2, c2 = _read_coder(args.coder2)
    problems = []
    if not (done1 and done2):
        problems.append("both workbooks must be marked completed = TRUE on the 'Coder' sheet")
    if not id1 or not id2 or id1 == id2:
        problems.append("coder_id must be filled and different in the two workbooks")
    if set(c1) != set(c2):
        problems.append("the two workbooks do not list the same case_id values")
    if problems:
        raise SystemExit("cannot merge:\n- " + "\n- ".join(problems))

    gold_rows, disagreements = [], []
    for cid in sorted(c1):
        a, b = c1[cid], c2[cid]
        sa, sb = _bool(a["in_scope"], f"{id1} {cid}"), _bool(b["in_scope"], f"{id2} {cid}")
        ra, rb = str(a["route"] or "").strip(), str(b["route"] or "").strip()
        for who, rec, scope in ((id1, a, sa), (id2, b, sb)):
            if scope and (rec["route"] in (None, "") or _num(rec["r"]) is None or _num(rec["n"]) is None):
                raise SystemExit(f"{who} {cid}: an in-scope case needs route, r and n")
            if rec["minutes"] in (None, ""):
                raise SystemExit(f"{who} {cid}: minutes is required")
        fields = []
        if sa != sb:
            fields.append("in_scope")
        if sa and sb:
            if ra != rb:
                fields.append("route")
            if abs(_num(a["r"]) - _num(b["r"])) > args.r_tolerance:
                fields.append("r")
            if int(_num(a["n"])) != int(_num(b["n"])):
                fields.append("n")
            for key in ("doi_measure", "performance_measure"):
                if str(a[key] or "").strip() != str(b[key] or "").strip():
                    fields.append(key)
        agreed = not fields
        row = {
            "case_id": cid, "paper_id": a["paper_id"], "effect_id": a["effect_id"], "source_locator": a["source_locator"],
            "coder1_in_scope": "TRUE" if sa else "FALSE", "coder2_in_scope": "TRUE" if sb else "FALSE",
            "coder1_route": ra if sa else "", "coder2_route": rb if sb else "",
            "manual_minutes": _num(a["minutes"]),
            "gold_in_scope": "", "gold_route": "", "gold_r": "", "gold_n": "",
            "gold_doi_measure": "", "gold_performance_measure": "", "adjudication_notes": "",
        }
        if agreed:
            row["gold_in_scope"] = "TRUE" if sa else "FALSE"
            if sa:
                row.update(gold_route=ra, gold_r=f"{_num(a['r']):.3f}", gold_n=int(_num(a["n"])),
                           gold_doi_measure=str(a["doi_measure"] or "").strip(),
                           gold_performance_measure=str(a["performance_measure"] or "").strip())
            row["adjudication_notes"] = "agreed"
        else:
            disagreements.append({
                "case_id": cid, "fields": ";".join(fields),
                **{f"{id1}_{k}": a[k] for k in ("in_scope", "route", "r", "n", "doi_measure", "performance_measure", "notes")},
                **{f"{id2}_{k}": b[k] for k in ("in_scope", "route", "r", "n", "doi_measure", "performance_measure", "notes")},
            })
        gold_rows.append(row)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    gold_out, dis_out = args.out_dir / "gold_standard_draft.csv", args.out_dir / "disagreements.csv"
    for p in (gold_out, dis_out):
        if p.exists():
            raise SystemExit(f"{p} exists; outputs are never overwritten")
    with gold_out.open("w", encoding="utf-8", newline="") as h:
        w = csv.DictWriter(h, fieldnames=GOLD_COLUMNS)
        w.writeheader()
        w.writerows(gold_rows)
    dis_cols = ["case_id", "fields"] + [f"{i}_{k}" for i in (id1, id2)
                                        for k in ("in_scope", "route", "r", "n", "doi_measure", "performance_measure", "notes")]
    with dis_out.open("w", encoding="utf-8", newline="") as h:
        w = csv.DictWriter(h, fieldnames=dis_cols)
        w.writeheader()
        w.writerows(disagreements)
    n = len(gold_rows)
    agree_scope = sum(r["coder1_in_scope"] == r["coder2_in_scope"] for r in gold_rows)
    print(f"merged {n} cases: {n - len(disagreements)} agreed, {len(disagreements)} to adjudicate "
          f"(raw scope agreement {agree_scope}/{n}); wrote {gold_out} and {dis_out}")
    print("Adjudicate each row of disagreements.csv against the full text, fill gold_* and adjudication_notes "
          "in a COPY named validation/gold_standard.csv; analyze_validation.py computes kappa from coder1_/coder2_ columns.")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("enumerate")
    e.add_argument("--lock", type=Path)
    e.add_argument("--sampling", type=Path, default=SAMPLING)
    e.add_argument("--rows-per-paper", type=int, default=4)
    e.add_argument("--out", type=Path, required=True)
    c = sub.add_parser("coders")
    c.add_argument("--enumeration", type=Path, required=True)
    c.add_argument("--out-dir", type=Path, required=True)
    m = sub.add_parser("merge")
    m.add_argument("--coder1", type=Path, required=True)
    m.add_argument("--coder2", type=Path, required=True)
    m.add_argument("--out-dir", type=Path, required=True)
    m.add_argument("--r-tolerance", type=float, default=0.005)
    args = ap.parse_args(argv)
    if args.cmd == "enumerate" and args.lock is None:
        print("note: no --lock given; the workbook lists the PROVISIONAL frame (rehearsal only)")
    return {"enumerate": cmd_enumerate, "coders": cmd_coders, "merge": cmd_merge}[args.cmd](args)


if __name__ == "__main__":
    raise SystemExit(main())
